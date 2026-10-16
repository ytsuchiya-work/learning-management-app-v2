"""生徒成績管理システム — FastAPI backend.

Serves JSON APIs for the React frontend plus the built static assets.
Everything runs in a single process (frontend + backend) to avoid the
Databricks Apps reverse-proxy CORS restriction.
"""
import os
import re
import json
import logging
from functools import lru_cache
from typing import Any, Optional

import pandas as pd
from fastapi import FastAPI, Request, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from databricks.sdk import WorkspaceClient

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("lms")

# ──────────────────────────────────────────────
# Configuration
# ──────────────────────────────────────────────
GENIE_SPACE_ID = os.environ.get("GENIE_SPACE_ID", "01f1c1250315122cac0d9bd97f5ba170")
WAREHOUSE_ID = os.environ.get("WAREHOUSE_ID", "e351c2d1b16eae95")
CATALOG = os.environ.get("CATALOG_NAME", "classic_stable_ytcy_catalog")
SCHEMA = os.environ.get("SCHEMA_NAME", "learning_management")
SERVING_ENDPOINT = os.environ.get("SERVING_ENDPOINT", "databricks-claude-sonnet-5-5")
# Unity Catalog governed model service (queried via the AI Gateway, not the
# pay-per-token FM endpoints which are disabled on this workspace).
MODEL_SERVICE = os.environ.get(
    "MODEL_SERVICE", "classic_stable_ytcy_catalog.learning_management.my_service_opus_5_5"
)
AI_GATEWAY_ROUTE = "/ai-gateway/mlflow/v1/chat/completions"

app = FastAPI(title="生徒成績管理システム")


@lru_cache(maxsize=1)
def get_workspace_client() -> WorkspaceClient:
    return WorkspaceClient()


# ──────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────
def _code(val: str) -> str:
    """Validate a classroom/student code is a safe numeric token."""
    if val is None or not re.fullmatch(r"\d{1,10}", str(val)):
        raise HTTPException(status_code=400, detail=f"invalid code: {val!r}")
    return str(val)


def run_sql(query: str) -> pd.DataFrame:
    w = get_workspace_client()
    resp = w.statement_execution.execute_statement(
        warehouse_id=WAREHOUSE_ID,
        statement=query,
        catalog=CATALOG,
        schema=SCHEMA,
        wait_timeout="50s",
    )
    state = resp.status.state.value if resp.status and resp.status.state else None
    if state == "SUCCEEDED":
        cols = [c.name for c in resp.manifest.schema.columns]
        rows = resp.result.data_array if resp.result and resp.result.data_array else []
        return pd.DataFrame(rows, columns=cols)
    err = resp.status.error if resp.status else "unknown error"
    raise HTTPException(status_code=500, detail=f"SQL error: {err}")


_KAISU_RE = re.compile(r"(\d+)")


def _kaisu_num(s: str) -> int:
    m = _KAISU_RE.search(str(s or ""))
    return int(m.group(1)) if m else 0


# オープン模試の実施月。データには「回数」しかないため、提案資料【例】の表記
# （第1回=4月・第3回=8月・第7回=1月）に合わせて変換する。対応の無い回は「第N回」のまま。
KAISU_MONTH = {1: "4月", 3: "8月", 7: "1月"}


def _period(grade: Any, kaisu: Any) -> str:
    m = KAISU_MONTH.get(_kaisu_num(kaisu))
    return f"{grade} {m or kaisu}"


def _ctest_series(ct: pd.DataFrame) -> list[dict[str, Any]]:
    """Cテスト rows → chronological points labelled by 学年＋実施月."""
    if ct.empty:
        return []
    ct = ct.copy()
    for col in SUBJECTS + ["5科合計", "5科偏差値"]:
        ct[col] = pd.to_numeric(ct[col], errors="coerce")
    ct["_order"] = ct.apply(lambda r: (str(r["受験年度"]), _kaisu_num(r["回数"])), axis=1)
    out = []
    for _, r in ct.sort_values("_order").iterrows():
        pt: dict[str, Any] = {"label": _period(r["学年"], r["回数"]), "year": str(r["受験年度"])}
        for s in SUBJECTS:
            pt[s] = None if pd.isna(r[s]) else int(r[s])
        pt["合計"] = None if pd.isna(r["5科合計"]) else int(r["5科合計"])
        pt["偏差値"] = None if pd.isna(r["5科偏差値"]) else int(r["5科偏差値"])
        out.append(pt)
    return out


def _pcs_units_sql(cc: str, sc: str) -> str:
    """PCS 単元別の集計。同じ設問が大単元名・単元名の表記ゆれ（1次関数/一次関数 等）の
    組み合わせごとに重複登録されているため、単元は (教科, 大単元NO, 連番) で識別し、
    設問は重複を除いて数える。単元名は表記ゆれのうち1つを代表として使う。"""
    src = f"{CATALOG}.{SCHEMA}.pcs_results WHERE `教室コード`={cc} AND `生徒コード`={sc}"
    return f"""
        WITH q AS (SELECT DISTINCT `教科`,`大単元NO`,`連番`,`実施回数`,`問題NO`,`正答率` FROM {src}),
             n AS (SELECT `教科`,`大単元NO`,`連番`, min(`大単元名`) `大単元名`, min(`単元名`) `単元名`
                   FROM {src} GROUP BY `教科`,`大単元NO`,`連番`)
        SELECT n.`教科`, n.`大単元名`, n.`単元名`,
               count(*) total,
               sum(CASE WHEN q.`正答率`='100.00%' THEN 1 ELSE 0 END) correct,
               sum(CASE WHEN q.`正答率` NOT IN ('100.00%','0.00%') THEN 1 ELSE 0 END) partial,
               sum(CASE WHEN q.`正答率`='0.00%' THEN 1 ELSE 0 END) incorrect
        FROM q JOIN n ON q.`教科`=n.`教科` AND q.`大単元NO`=n.`大単元NO` AND q.`連番`=n.`連番`
        GROUP BY n.`教科`, n.`大単元NO`, n.`連番`, n.`大単元名`, n.`単元名`
        ORDER BY n.`教科`, n.`大単元NO`, n.`連番`
    """


def _pcs_status(total: int, correct: int) -> str:
    if total > 0 and correct == total:
        return "定着"
    if total > 0 and correct >= total * 0.5:
        return "一部定着"
    return "未定着"


def _unit_name(major: Any, unit: Any) -> str:
    """「大単元（単元）」— the way the proposal sheet names units."""
    major, unit = str(major or "").strip(), str(unit or "").strip()
    return major if not unit or unit == major else f"{major}（{unit}）"


# ──────────────────────────────────────────────
# Identity
# ──────────────────────────────────────────────
@app.get("/api/whoami")
def whoami(request: Request):
    email = (
        request.headers.get("x-forwarded-email")
        or request.headers.get("x-forwarded-user")
        or ""
    )
    return {"email": email, "sp_note": "データ照会はアプリのサービスプリンシパルとして実行されます"}


# ──────────────────────────────────────────────
# Students
# ──────────────────────────────────────────────
@app.get("/api/students")
def students():
    df = run_sql(
        f"SELECT classroom_code, student_code FROM {CATALOG}.{SCHEMA}.students "
        f"ORDER BY classroom_code, student_code"
    )
    by_room: dict[str, list[str]] = {}
    for _, r in df.iterrows():
        by_room.setdefault(str(r["classroom_code"]), []).append(str(r["student_code"]))
    classrooms = sorted(by_room.keys())
    return {"classrooms": classrooms, "by_classroom": by_room}


# ──────────────────────────────────────────────
# Dashboard
# ──────────────────────────────────────────────
SUBJECTS = ["英語", "数学", "国語", "理科", "社会"]


@app.get("/api/dashboard")
def dashboard(classroom: str, student: str, mode: str = "管理者"):
    cc, sc = _code(classroom), _code(student)

    counts = {}
    for key, tbl, where in [
        ("ctest", "c_test_results", f"`教室コード`={cc} AND `生徒コード`={sc}"),
        ("pcs", "pcs_results", f"`教室コード`={cc} AND `生徒コード`={sc}"),
        ("ets", "ets_metadata", f"classroom_code='{cc}' AND student_code='{sc}'"),
        ("sss", "sss_metadata", f"classroom_code='{cc}' AND student_code='{sc}'"),
    ]:
        d = run_sql(f"SELECT count(*) cnt FROM {CATALOG}.{SCHEMA}.{tbl} WHERE {where}")
        counts[key] = int(d.iloc[0]["cnt"]) if not d.empty else 0

    # ── C-test score trend ──
    ctest_series: list[dict[str, Any]] = []
    deviation_latest: dict[str, Any] = {}
    if counts["ctest"] > 0:
        ct = run_sql(f"""
            SELECT `受験年度`,`学年`,`回数`,`英語`,`数学`,`国語`,`理科`,`社会`,`5科合計`,`5科偏差値`
            FROM {CATALOG}.{SCHEMA}.c_test_results
            WHERE `教室コード`={cc} AND `生徒コード`={sc}
        """)
        ctest_series = _ctest_series(ct)
        if ctest_series:
            last = ctest_series[-1]
            deviation_latest = {"偏差値": last["偏差値"], "合計": last["合計"], "label": last["label"]}

    # ── PCS mastery ──
    pcs_units: list[dict[str, Any]] = []
    pcs_summary = {"定着": 0, "一部定着": 0, "未定着": 0}
    if counts["pcs"] > 0:
        pcs = run_sql(_pcs_units_sql(cc, sc))
        for col in ["total", "correct", "partial", "incorrect"]:
            pcs[col] = pd.to_numeric(pcs[col], errors="coerce").fillna(0).astype(int)
        for _, r in pcs.iterrows():
            total = int(r["total"])
            correct = int(r["correct"])
            status = _pcs_status(total, correct)
            pcs_summary[status] += 1
            pcs_units.append({
                "教科": r["教科"], "大単元名": r["大単元名"], "単元名": r["単元名"],
                "問題数": total, "正解": correct, "部分正解": int(r["partial"]),
                "不正解": int(r["incorrect"]), "定着度": status,
            })

    # ── Classroom overview (admin only) ──
    overview: list[dict[str, Any]] = []
    if mode == "管理者":
        ov = run_sql(f"""
            SELECT CAST(`生徒コード` AS STRING) student_code,
                   MAX(`英語`) max_eng, MAX(`数学`) max_math, MAX(`国語`) max_jpn,
                   MAX(`理科`) max_sci, MAX(`社会`) max_soc, MAX(`5科合計`) max_total
            FROM {CATALOG}.{SCHEMA}.c_test_results
            WHERE `教室コード`={cc}
            GROUP BY `生徒コード`
            ORDER BY max_total DESC
        """)
        for col in ["max_eng", "max_math", "max_jpn", "max_sci", "max_soc", "max_total"]:
            ov[col] = pd.to_numeric(ov[col], errors="coerce")
        for _, r in ov.iterrows():
            overview.append({
                "生徒コード": r["student_code"],
                "英語": None if pd.isna(r["max_eng"]) else int(r["max_eng"]),
                "数学": None if pd.isna(r["max_math"]) else int(r["max_math"]),
                "国語": None if pd.isna(r["max_jpn"]) else int(r["max_jpn"]),
                "理科": None if pd.isna(r["max_sci"]) else int(r["max_sci"]),
                "社会": None if pd.isna(r["max_soc"]) else int(r["max_soc"]),
                "5科合計": None if pd.isna(r["max_total"]) else int(r["max_total"]),
                "is_current": r["student_code"] == sc,
            })

    return {
        "counts": counts,
        "ctest_series": ctest_series,
        "deviation_latest": deviation_latest,
        "pcs_units": pcs_units,
        "pcs_summary": pcs_summary,
        "overview": overview,
        "subjects": SUBJECTS,
    }


# ──────────────────────────────────────────────
# Genie chat
# ──────────────────────────────────────────────
class GenieAsk(BaseModel):
    message: str
    conversation_id: Optional[str] = None


@app.post("/api/genie/ask")
def genie_ask(body: GenieAsk):
    w = get_workspace_client()
    try:
        if not body.conversation_id:
            resp = w.genie.start_conversation_and_wait(
                space_id=GENIE_SPACE_ID, content=body.message
            )
            conversation_id = resp.conversation_id
        else:
            resp = w.genie.create_message_and_wait(
                space_id=GENIE_SPACE_ID,
                conversation_id=body.conversation_id,
                content=body.message,
            )
            conversation_id = body.conversation_id

        answer_text = ""
        sql_query = None
        columns: list[str] = []
        rows: list[list[Any]] = []
        if resp.attachments:
            for att in resp.attachments:
                if att.text and att.text.content:
                    answer_text += att.text.content + "\n"
                if att.query and att.query.query:
                    sql_query = att.query.query
                    try:
                        df = run_sql(att.query.query)
                        columns = list(df.columns)
                        rows = df.values.tolist()
                    except Exception as e:
                        log.warning("genie sql failed: %s", e)

        if not answer_text:
            answer_text = "回答を取得できませんでした。別の質問をお試しください。"

        return {
            "answer": answer_text.strip(),
            "conversation_id": conversation_id,
            "sql": sql_query,
            "columns": columns,
            "rows": rows,
        }
    except Exception as e:
        log.exception("genie error")
        raise HTTPException(status_code=500, detail=str(e))


# ──────────────────────────────────────────────
# Parent conference material (LLM)
# ──────────────────────────────────────────────
class MendanReq(BaseModel):
    classroom: str
    student: str
    tone: str = "① 励まし重視"
    target_school: str = ""
    current_dev: str = ""
    target_dev: str = ""
    mode: str = "structured"  # "structured" (定型: 講習提案書) | "free" (自由フォーマット)
    ref_sss: list[str] = []   # 参考にする他の先生の講習SSS（Volume file_path）
    # 提案書の見出し・裏面に載せる、データに無い情報（教室長が入力）
    student_name: str = ""
    course: str = ""          # 学科・コース
    judgement: str = ""       # 判定（例: B判定）
    school_tests: dict[str, dict[str, str]] = {}  # 定期テスト {教科: {"score": 点, "target": 目標}}
    koma_individual: dict[str, str] = {}          # 個別の提案コマ数 {教科: コマ数}
    koma_group: dict[str, str] = {}               # 集団・その他の提案コマ数 {教科: コマ数}
    koma_other: str = ""                          # 5教科以外の提案内容


def _chat_llm(messages: list[dict], max_tokens: int = 8000, reasoning_effort: str | None = None) -> str:
    """Query the Unity Catalog governed model service via the AI Gateway
    (OpenAI-compatible). Returns the assistant message text.

    Uses `requests` directly instead of `api_client.do`: the SDK client has a
    60s read timeout and silently *re-sends* the request on timeout (retrying
    for up to 5 min), so long Opus generations never finish and get billed
    several times. Here: one request, 10 min timeout, no retry."""
    import requests

    cfg = get_workspace_client().config
    headers = {"Content-Type": "application/json", **cfg.authenticate()}
    body: dict[str, Any] = {"model": MODEL_SERVICE, "messages": messages, "max_tokens": max_tokens}
    if reasoning_effort:
        body["reasoning_effort"] = reasoning_effort
    r = requests.post(f"{cfg.host.rstrip('/')}{AI_GATEWAY_ROUTE}", json=body, headers=headers, timeout=600)
    if r.status_code >= 400:
        raise RuntimeError(f"model service error {r.status_code}: {r.text[:300]}")
    resp = r.json()
    choices = resp.get("choices") or []
    if not choices:
        return ""
    return _content_to_text(choices[0].get("message", {}).get("content"))


def _content_to_text(content: Any) -> str:
    """Serving responses may return `content` as a plain string or, for Claude
    on the Foundation Model API, a list of content blocks. Normalise to text."""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = []
        for block in content:
            if isinstance(block, dict):
                parts.append(block.get("text") or block.get("content") or "")
            else:
                parts.append(getattr(block, "text", "") or str(block))
        return "".join(parts)
    return str(content or "")


def _gather_student_data(cc: str, sc: str) -> dict[str, pd.DataFrame]:
    data = {}
    data["ctest"] = run_sql(f"""
        SELECT `学年`,`回数`,`受験年度`,`英語`,`数学`,`国語`,`理科`,`社会`,`5科合計`,`5科偏差値`
        FROM {CATALOG}.{SCHEMA}.c_test_results
        WHERE `教室コード`={cc} AND `生徒コード`={sc}
        ORDER BY `受験年度`,`回数`
    """)
    data["pcs"] = run_sql(_pcs_units_sql(cc, sc))
    data["ets"] = run_sql(f"""
        SELECT file_name, file_path FROM {CATALOG}.{SCHEMA}.ets_metadata
        WHERE classroom_code='{cc}' AND student_code='{sc}'
    """)
    data["sss"] = run_sql(f"""
        SELECT file_name, file_path FROM {CATALOG}.{SCHEMA}.sss_metadata
        WHERE classroom_code='{cc}' AND student_code='{sc}'
    """)
    return data


@app.get("/api/mendan/data")
def mendan_data(classroom: str, student: str):
    cc, sc = _code(classroom), _code(student)
    data = _gather_student_data(cc, sc)
    return {
        "ctest": len(data["ctest"]),
        "pcs": len(data["pcs"]),
        "ets": len(data["ets"]),
        "sss": len(data["sss"]),
    }


@lru_cache(maxsize=128)
def _parse_ets_text(file_path: str) -> str:
    """Extract text from an ETS 「やる気度診断テスト」 PDF via ai_parse_document.
    Cached per file (content is static). Returns '' on any failure so the
    conference-material generation degrades gracefully to quantitative signals."""
    if not re.fullmatch(r"[\w/.\-一-龥ぁ-んァ-ヶ]+", file_path or ""):
        return ""
    try:
        w = get_workspace_client()
        resp = w.statement_execution.execute_statement(
            warehouse_id=WAREHOUSE_ID,
            statement=(
                "SELECT try_cast(ai_parse_document(content) AS STRING) parsed "
                f"FROM READ_FILES('{file_path}', format => 'binaryFile')"
            ),
            wait_timeout="50s",
        )
        if not (resp.status and resp.status.state and resp.status.state.value == "SUCCEEDED"):
            return ""
        if not (resp.result and resp.result.data_array):
            return ""
        parsed = resp.result.data_array[0][0]
        doc = json.loads(parsed)
        wanted = {"title", "section_header", "text", "table"}
        seen, parts = set(), []
        for el in doc.get("document", {}).get("elements", []):
            if el.get("type") not in wanted:
                continue
            txt = (el.get("content") or "").strip()
            if txt and txt not in seen and len(txt) > 1:
                seen.add(txt)
                parts.append(txt)
        # 2ページ目の「こうすればもっと成績は上がる」まで届くよう広めに取る
        return "\n".join(parts)[:7000]
    except Exception as e:  # noqa: BLE001
        log.warning("ETS parse failed for %s: %s", file_path, e)
        return ""


def _right_of(ws, label: str) -> str:
    """Value of the first non-empty cell to the right of a cell whose text == label."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() == label:
                for col in range(c.column + 1, c.column + 6):
                    v = ws.cell(row=c.row, column=col).value
                    if v not in (None, ""):
                        return str(v).strip()
                return ""
    return ""


def _below(ws, header: str, skip: str = "保護者面談") -> str:
    """First meaningful text below a section header (same column)."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and header in str(c.value):
                for r in range(c.row + 1, c.row + 6):
                    v = ws.cell(row=r, column=c.column).value
                    if v not in (None, "") and skip not in str(v):
                        return str(v).strip()
                return ""
    return ""


@lru_cache(maxsize=128)
def _load_sss(file_path: str) -> tuple[str, dict]:
    """Read a 講習SSS (.xlsx) — the per-manager (室長) conference sheet with
    入会ヒアリング・目標設定・室長提案 — via the Files API + openpyxl.
    Returns (full text, profile). Cached per file; ('', {}) on failure."""
    if not file_path or not file_path.startswith("/Volumes/"):
        return "", {}
    try:
        import io
        import openpyxl

        w = get_workspace_client()
        resp = w.files.download(file_path)
        raw = resp.contents.read() if hasattr(resp.contents, "read") else resp.contents
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        ws = wb.worksheets[0]
        head = str(ws["B2"].value or "")
        grade = re.search(r"\[(.+?)\]", head)
        profile = {
            "school": str(ws["G2"].value or "").strip(),
            "term": ws.title,
            "grade": grade.group(1) if grade else "",
            "teachers": _right_of(ws, "担当講師"),
            "target": _right_of(ws, "志望校・進路"),
            "compat": _right_of(ws, "相性の合うタイプ"),
            "course": _right_of(ws, "コース・週回数"),
            "proposal": _below(ws, "■提案内容（室長）"),
            "report": _below(ws, "■報告内容"),
        }
        lines, seen = [], set()
        for sh in wb.worksheets:
            for row in sh.iter_rows(values_only=True):
                vals = [str(c).strip() for c in row if c not in (None, "")]
                if not vals:
                    continue
                line = " ｜ ".join(vals)
                if line not in seen and len(line) > 1:
                    seen.add(line)
                    lines.append(line)
            if sum(len(x) for x in lines) > 3200:
                break
        return "\n".join(lines)[:3200], profile
    except Exception as e:  # noqa: BLE001
        log.warning("SSS parse failed for %s: %s", file_path, e)
        return "", {}


def _parse_sss_text(file_path: str) -> str:
    return _load_sss(file_path)[0]


def _sss_label(p: dict) -> str:
    return f"{p.get('school', '')} {p.get('term', '')}（担当: {p.get('teachers') or '—'}）".strip()


_EMPTY_TARGETS = {"", "なし", "未定", "無し", "-", "—"}


@app.get("/api/mendan/prefill")
def mendan_prefill(classroom: str, student: str):
    """Pre-fill 志望校 / 現状偏差値 / 目標偏差値 from the student's own data.
    Each value carries the source it came from; '' means "not found → 手入力"."""
    cc, sc = _code(classroom), _code(student)
    out = {"target_school": "", "current_dev": "", "target_dev": "", "sources": {}}

    sss = run_sql(
        f"SELECT file_path FROM {CATALOG}.{SCHEMA}.sss_metadata "
        f"WHERE classroom_code='{cc}' AND student_code='{sc}'"
    )
    sss_text = ""
    if not sss.empty:
        sss_text, prof = _load_sss(str(sss.iloc[0]["file_path"]))
        tgt = (prof.get("target") or "").strip()
        if tgt not in _EMPTY_TARGETS:
            out["target_school"] = tgt
            out["sources"]["target_school"] = f"講習SSS（志望校・進路）{_sss_label(prof)}"

    ct = run_sql(f"""
        SELECT `受験年度`,`学年`,`回数`,`5科偏差値` FROM {CATALOG}.{SCHEMA}.c_test_results
        WHERE `教室コード`={cc} AND `生徒コード`={sc} AND `5科偏差値` IS NOT NULL
    """)
    if not ct.empty:
        ct["_o"] = ct.apply(lambda r: (str(r["受験年度"]), _kaisu_num(r["回数"])), axis=1)
        last = ct.sort_values("_o").iloc[-1]
        out["current_dev"] = str(last["5科偏差値"])
        out["sources"]["current_dev"] = f"c_test_results（{last['学年']} {last['回数']} の5科偏差値）"

    m = re.search(r"目標偏差値\D{0,4}(\d{2})", sss_text or "")
    if m:
        out["target_dev"] = m.group(1)
        out["sources"]["target_dev"] = "講習SSS（目標偏差値）"
    return out


@app.get("/api/sss/list")
def sss_list():
    """All 講習SSS sheets across classrooms, with a profile per sheet so the user
    can pick which other teacher's sheet to use as a reference."""
    df = run_sql(
        f"SELECT classroom_code, student_code, file_name, file_path "
        f"FROM {CATALOG}.{SCHEMA}.sss_metadata ORDER BY classroom_code, student_code"
    )
    rows = [r for _, r in df.iterrows()]
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=8) as ex:
        profiles = list(ex.map(lambda r: _load_sss(str(r["file_path"]))[1], rows))
    out = []
    for r, p in zip(rows, profiles):
        out.append({
            "classroom": str(r["classroom_code"]), "student": str(r["student_code"]),
            "file_name": r["file_name"], "file_path": r["file_path"],
            "label": _sss_label(p), **p,
        })
    return {"items": out}


def _behavior_signals(ctest: pd.DataFrame, pcs: pd.DataFrame) -> str:
    """Quantitative learning-characteristic signals derived from Cテスト/PCS,
    used to ground the 学習特性・行動評価 section even without ETS content."""
    sig: list[str] = []
    if not ctest.empty:
        c = ctest.copy()
        for col in SUBJECTS + ["5科合計", "5科偏差値"]:
            if col in c:
                c[col] = pd.to_numeric(c[col], errors="coerce")
        for s in SUBJECTS:
            vals = [int(v) for v in c[s].dropna().tolist()]
            if len(vals) >= 2:
                rng = max(vals) - min(vals)
                delta = vals[-1] - vals[0]
                trend = "上昇" if delta >= 8 else ("下降" if delta <= -8 else "横ばい")
                stab = "安定" if rng <= 15 else ("変動大" if rng >= 30 else "やや変動")
                sig.append(f"{s}: 初回{vals[0]}→直近{vals[-1]}（{trend}・{stab}, 範囲{min(vals)}-{max(vals)}）")
        tot = [int(v) for v in c["5科合計"].dropna().tolist()]
        if len(tot) >= 2:
            d = tot[-1] - tot[0]
            sig.append(f"5科合計: {tot[0]}→{tot[-1]}（{'上昇' if d > 0 else '下降' if d < 0 else '横ばい'}）")
        if "5科偏差値" in c:
            dv = [int(v) for v in c["5科偏差値"].dropna().tolist()]
            if dv:
                sig.append(f"5科偏差値: 初回{dv[0]}→直近{dv[-1]}")
    if not pcs.empty:
        p = pcs.copy()
        for col in ["total", "correct"]:
            p[col] = pd.to_numeric(p[col], errors="coerce").fillna(0).astype(int)

        def _st(r):
            t, co = int(r["total"]), int(r["correct"])
            return "定着" if t > 0 and co == t else ("一部定着" if t > 0 and co >= t * 0.5 else "未定着")

        p["st"] = p.apply(_st, axis=1)
        for subj, grp in p.groupby("教科"):
            tot = len(grp)
            m = int((grp["st"] == "定着").sum())
            pt = int((grp["st"] == "一部定着").sum())
            un = int((grp["st"] == "未定着").sum())
            sig.append(f"PCS {subj}: 全{tot}単元中 定着{m}・一部{pt}・未定着{un}（定着率{round(m / tot * 100) if tot else 0}%）")
        und = p[p["st"] == "未定着"]["大単元名"].value_counts().head(8)
        if len(und):
            sig.append("未定着が多い大単元: " + "、".join(f"{k}({v})" for k, v in und.items()))
    return "\n".join(sig) if sig else "（定量シグナルを算出できる成績データがありません）"


def _field_sources(data: dict, ets_used: bool, sss_used: bool, ref_labels: list[str] | None = None,
                   req: Optional["MendanReq"] = None) -> dict[str, list[str]]:
    """Deterministic provenance per output field so each section shows which
    governed table / document it is grounded in (addresses traceability).
    Shown on screen only; the printed proposal uses parent-friendly notes instead."""
    ct = not data["ctest"].empty
    pc = not data["pcs"].empty
    teiki = bool(req and _school_tests_text(req))
    sss_src = ["sss_metadata → 講習SSS（本人・室長作成の面談シート）"] if sss_used else []
    ref_src = [f"参考SSS → {lb}" for lb in (ref_labels or [])]
    teiki_src = ["面談入力（定期テスト）"] if teiki else []
    s: dict[str, list[str]] = {}
    s["student_goal_text"] = ["面談入力（志望校・コース）"] + (["c_test_results（オープン模試）"] if ct else []) + sss_src
    s["ctest_summary"] = (["c_test_results（オープン模試 得点推移）"] if ct else ["（成績データなし）"]) + teiki_src
    s["target_gap"] = (["c_test_results（5科偏差値）"] if ct else []) + ["面談入力（志望校・目標偏差値）"] \
        + teiki_src + (["pcs_results（定着している単元）"] if pc else []) + sss_src
    ev = []
    if ets_used:
        ev.append("ets_metadata → ETSやる気度診断テスト（学習特性・行動評価）")
    if ct:
        ev.append("c_test_results（得点の安定性・推移）")
    if pc:
        ev.append("pcs_results（単元定着の偏り）")
    ev += sss_src
    s["ets_eval"] = ev or ["（根拠データなし）"]
    s["english_analysis"] = (["c_test_results（英語得点）"] if ct else []) + (["pcs_results（英語単元）"] if pc else [])
    s["math_analysis"] = (["c_test_results（数学得点）"] if ct else []) + (["pcs_results（数学単元）"] if pc else [])
    s["other_subjects"] = ["c_test_results（国語・理科・社会の得点）"] if ct else ["（成績データなし）"]
    s["remarks"] = ["教室長が記入"]
    s["koma_units"] = ["面談入力（提案コマ数）"] + (["pcs_results（未定着・一部定着の単元）"] if pc else [])
    # Reference sheets shape the goal-setting and proposal sections (書き方・提案の型)
    for k in ("student_goal_text", "target_gap", "english_analysis", "math_analysis", "koma_units"):
        s[k] = s[k] + ref_src
    s["data_notes"] = ["システム内部（データ欠落・根拠メモ）"]
    return s


def _fallback_ets_eval(signals: str, ets_text: str) -> str:
    """Deterministic 学習特性・行動評価 built from quantitative signals, used when
    the LLM leaves the section blank. Always data-grounded, never empty."""
    lines = ["【成績データに基づく学習特性】"]
    for ln in signals.splitlines():
        if ln.strip() and not ln.startswith("（"):
            lines.append("・" + ln.strip())
    if ets_text:
        lines.append("※ ETSやる気度診断テストの結果も参照（詳細は診断PDFを担当者が保管）。")
    return "\n".join(lines)


_ETS_DATE_RE = re.compile(r"(\d{4})\s*年\s*0?(\d{1,2})\s*月")


def _ets_date(ets_text: str) -> str:
    """Date of the ETS diagnosis (e.g. '2025年4月') for the「※最新のETSは…時点」note."""
    m = _ETS_DATE_RE.search(ets_text or "")
    return f"{m.group(1)}年{int(m.group(2))}月" if m else ""


def _school_tests_text(req: "MendanReq") -> str:
    rows = []
    for s in SUBJECTS:
        t = (req.school_tests or {}).get(s) or {}
        score, target = str(t.get("score") or "").strip(), str(t.get("target") or "").strip()
        if score or target:
            rows.append(f"{s}: {score or '—'}点" + (f"（目標 {target}点）" if target else ""))
    return "\n".join(rows)


def _koma_int(v: Any) -> int:
    m = re.search(r"\d+", str(v or ""))
    return int(m.group(0)) if m else 0


def _build_context(data: dict, signals: str, ets_text: str, sss_text: str, ref_text: str = "",
                   req: Optional["MendanReq"] = None) -> str:
    """Shared data-body (A〜H) used by both the structured and free prompts."""
    sections = []
    series = _ctest_series(data["ctest"])
    if series:
        lines = ["時期 | 英語 | 数学 | 国語 | 理科 | 社会 | 5科合計 | 5科偏差値"]
        for p in series:
            vals = [p[s] for s in SUBJECTS] + [p["合計"], p["偏差値"]]
            lines.append(p["label"] + " | " + " | ".join("（未受験）" if v is None else str(v) for v in vals))
        sections.append("【A】オープン模試 教科別得点推移（時系列順）:\n" + "\n".join(lines))
    else:
        sections.append("【A】オープン模試: データなし")

    if not data["pcs"].empty:
        by_subj: dict[str, dict[str, list[str]]] = {}
        for _, r in data["pcs"].iterrows():
            total = int(r.get("total", 0) or 0)
            correct = int(r.get("correct", 0) or 0)
            st = _pcs_status(total, correct)
            subj = "数学" if "数学" in str(r["教科"]) else str(r["教科"])
            by_subj.setdefault(subj, {"定着": [], "一部定着": [], "未定着": []})[st].append(
                _unit_name(r["大単元名"], r["単元名"]))
        lines = []
        for subj, g in by_subj.items():
            lines.append(f"＜{subj}＞")
            for st in ("定着", "一部定着", "未定着"):
                lines.append(f"{st}: " + ("、".join(g[st]) if g[st] else "なし"))
        sections.append(
            "【B】単元別の学習記録（PCS）— 全単元。定着=全問正解、一部定着=半分以上正解、未定着=それ未満。"
            "単元名は「大単元（単元）」表記:\n" + "\n".join(lines))
    else:
        sections.append("【B】単元別の学習記録（PCS）: データなし")

    sections.append("【C】学習特性の定量シグナル（模試・単元の集計。本文に数値を転記しない）:\n" + signals)

    if ets_text:
        sections.append("【D】学習行動診断（ETS やる気度診断テスト）の抽出テキスト:\n" + ets_text)
    elif not data["ets"].empty:
        sections.append(f"【D】学習行動診断（ETS）: {len(data['ets'])} 件あり（本文抽出不可。【C】で代替）")
    else:
        sections.append("【D】学習行動診断（ETS）: なし（【C】で評価）")

    if sss_text:
        sections.append("【E】講習SSS 本人の面談シート（室長の入会ヒアリング・目標設定・提案内容）:\n" + sss_text)
    else:
        sections.append("【E】講習SSS（本人）: なし")

    if ref_text:
        sections.append(
            "【F】参考SSS 他の先生（室長）が作成した講習SSS — 書き方・目標設定・提案の型の参考用。"
            "※この生徒のデータではないため、数値・単元名・志望校・個人情報を転記してはいけない:\n" + ref_text
        )

    if req is not None:
        teiki = _school_tests_text(req)
        sections.append("【G】学校の定期テスト（面談入力）:\n" + (teiki or "入力なし（定期テストにはふれない）"))
        koma = [f"{s}: {_koma_int(req.koma_individual.get(s))}コマ" for s in SUBJECTS
                if _koma_int((req.koma_individual or {}).get(s)) > 0]
        sections.append("【H】個別の提案コマ数（面談入力。koma_units はこの合計に合わせて配分する）:\n"
                        + ("\n".join(koma) if koma else "入力なし（koma_units はすべて空配列）"))

    return "\n\n".join(sections)


def _header(cc: str, sc: str, req: MendanReq) -> str:
    return (
        f"■ トーン: {req.tone}\n"
        f"■ 生徒名: {req.student_name or '（未入力）'}\n"
        f"■ 志望校: {req.target_school or '（未入力）'} / 学科・コース: {req.course or '（未入力）'} / 判定: {req.judgement or '（未入力）'}\n"
        f"■ 現状偏差値: {req.current_dev or '（未入力）'} / 目標偏差値: {req.target_dev or '（未入力）'}\n"
        f"■ 教室コード: {cc} / 生徒コード: {sc}"
    )


_REF_RULE = (
    "- 【F】参考SSSがある場合は、その先生の目標設定の立て方・提案の組み立て（単元とコマ数の示し方）・"
    "保護者への伝え方を参考にして、この生徒の内容を同じ型で書くこと。【F】の数値・志望校・単元名・氏名は絶対に使わないこと"
)

# 保護者に配る紙で使う言葉（提案資料【例】に合わせる）
_PARENT_WORDING = """- 読み手は保護者。システム名・テーブル名・データの欠損や表記ゆれの話は本文に書かない（data_notes にだけ書く）
- 言い換え: PCS →「これまでの単元別の学習記録」、ETS・やる気度診断テスト →「学習行動診断」、Cテスト →「オープン模試」
- 正答率・定着率・単元数・問題数・正解数などの集計値は書かない。数値は模試・定期テストの点数と偏差値を要所だけ使う
- 模試の時期は【A】の「時期」を使い「中学2年生の4月のオープン模試」のように書く（「第N回」とは書かない）
- 単元名は【B】にあるものだけを「大単元（単元）」の形で書く。【B】に無い単元名を作らない"""


def _build_prompt(cc: str, sc: str, data: dict, req: MendanReq, signals: str, ets_text: str, sss_text: str, ref_text: str = "") -> str:
    context = _build_context(data, signals, ets_text, sss_text, ref_text, req)
    return f"""あなたは個別指導塾の教室長です。保護者面談で配布する「講習 提案書」の本文を作成してください。
提案書は両面で、表面が本文、裏面が提案コマ数の表とオープン模試の推移表・グラフです（裏面の表はシステムが作るので、本文に表は書かない）。

【最重要原則】
- データ(A〜H)に存在しない数値・単元名・事実を創作してはいけません
- 出典ラベルは本文に書かないこと（別途システムが付与します）

【書き方の型（すべての章で厳守）】
- いつも前向きな言い方でまとめる。下がった点や苦手な点も「一時的な揺り戻し」「あと一歩で仕上がる状態」「これから積み上げたいポイント」のように、次につながる表現にする（トーン: {req.tone}）
- 1〜6章（ctest_summary〜other_subjects）は、それぞれ「・」で始まる行をちょうど4行。1行は70〜85字（1〜2文）。行は改行（\\n）で区切る
- 4行の順番: 1行目=しっかりできている点 → 2行目=できている点の続き、またはあと一歩の点 → 3行目=これから伸ばしたい点 → 4行目=「今後は、」で始まる方針
{_PARENT_WORDING}

【章ごとの内容】
- student_goal_text（目標）: 「・」を付けない3段落（改行区切り、各70〜85字）。①志望校合格に向けて今取り組んでいること ②強化したい教科と、それで広がる志望校の選択肢 ③普段の学習姿勢の良い点と、それを生かした次の一歩。偏差値などの数値目標は書かない
- ctest_summary（1. 全体成績推移）: オープン模試の得点の推移（最高だった時期・直近・伸びている教科）。【G】の定期テストがあれば4行目でふれる
- target_gap（2. 志望校ギャップ・課題）: 1行目=志望校と「現状偏差値X から目標偏差値Y まで、Z.0の差」（差は小数1桁）。2行目=【G】の定期テストの目標との差（無ければ偏差値の差を縮める道筋）。3行目=単元別の学習記録で定着している単元。4行目=応用・実践演習で差を縮める方針
- ets_eval（3. 学習特性・行動評価）: 【D】に基づく。診断の力の名前は【D】の表記どおり「」で囲んで書く。1〜3行目は、診断の本文で「優れている」「素晴らしい」とされた力や、グラフで高い力を強みとして書く。診断で「あなたに必要なもの」「こうすればもっと成績は上がる」「よりやる気を高めるためには」とされた項目は強みとして書かず、4行目の「今後は、」で伸ばしどころとして書く。【D】が無ければ【C】の傾向から書く
- english_analysis（4. 英語：現状分析）/ math_analysis（5. 数学：現状分析）: 1行目=定着している単元（3つ程度）、2行目=一部定着の単元を「あと一歩で仕上がる状態」として、3行目=未定着の単元を「これから積み上げたいポイント」として、4行目=志望校に向けた今後の方針
- other_subjects（6. 他教科の所見）: 国語・理科・社会について、オープン模試の得点の推移から書く
- remarks（備考）: 空文字 ""（教室長が記入する欄）
- koma_units（裏面の提案コマ数）: 【H】でコマ数がある教科だけ、合計がそのコマ数と一致するように、その教科の未定着・一部定着の単元（【B】の「大単元（単元）」表記）に配分する。1教科3〜8単元、未定着を優先。コマ数が無い教科は空配列
- data_notes（担当者用メモ・印刷しない）: データの欠損・不整合・表記ゆれ・根拠。ここだけは率直に書いてよい
- 【E】講習SSS（室長の入会ヒアリング・目標設定・提案内容）がある場合は、student_goal_text と target_gap にその内容を反映すること
{_REF_RULE}

{_header(cc, sc, req)}

■ データ本体:
{context}

次のJSONだけを出力してください:
{{
  "student_goal_text": "1段落目\\n2段落目\\n3段落目",
  "ctest_summary": "・1行目\\n・2行目\\n・3行目\\n・4行目",
  "target_gap": "・…（4行）",
  "ets_eval": "・…（4行）",
  "english_analysis": "・…（4行）",
  "math_analysis": "・…（4行）",
  "other_subjects": "・…（4行）",
  "remarks": "",
  "koma_units": {{"英語": [{{"unit": "大単元（単元）", "koma": 0}}], "数学": [], "国語": [], "理科": [], "社会": []}},
  "data_notes": "担当者用メモ"
}}"""


def _build_free_prompt(cc: str, sc: str, data: dict, req: MendanReq, signals: str, ets_text: str, sss_text: str, ref_text: str = "") -> str:
    context = _build_context(data, signals, ets_text, sss_text, ref_text, req)
    return f"""あなたは個別指導塾の教室長です。
保護者面談で配布する提案資料の本文を、**Markdown形式の自由な構成**で作成してください。
決まった項目に縛られず、この生徒に最も伝わる見出し・順序・強調で構成して構いません。

【厳守事項】
- データ(A〜H)に存在しない数値・単元名を創作しないこと（示された事実のみ）
- 前向きな言い方でまとめること（トーン: {req.tone}）
{_PARENT_WORDING}
- 学習行動診断や講習SSS（室長ノウハウ）がある場合は積極的に活用すること
{_REF_RULE}
- 見出しは `##`、強調は `**`、箇条書きは `- ` を使い、読みやすく整えること

{_header(cc, sc, req)}

■ データ本体:
{context}

Markdown本文のみを出力してください（前置きや説明は不要）。"""


def _fix_koma_units(raw: Any, req: MendanReq, data: dict) -> dict[str, list[dict[str, Any]]]:
    """Make the per-unit コマ数 add up to what the 教室長 entered. Subjects without
    an entered total get no units; if the LLM gave none, spread the total over the
    weakest PCS units (未定着 first)."""
    raw = raw if isinstance(raw, dict) else {}
    weak: dict[str, list[str]] = {}
    for _, r in data["pcs"].iterrows():
        st = _pcs_status(int(r.get("total", 0) or 0), int(r.get("correct", 0) or 0))
        if st != "定着":
            subj = "数学" if "数学" in str(r["教科"]) else str(r["教科"])
            weak.setdefault(subj, []).append((0 if st == "未定着" else 1, _unit_name(r["大単元名"], r["単元名"])))
    out: dict[str, list[dict[str, Any]]] = {}
    for s in SUBJECTS:
        total = _koma_int((req.koma_individual or {}).get(s))
        items = []
        for it in raw.get(s) or []:
            if isinstance(it, dict) and str(it.get("unit") or "").strip():
                items.append({"unit": str(it["unit"]).strip(), "koma": max(1, _koma_int(it.get("koma")))})
        if total <= 0:
            out[s] = []
            continue
        if not items:
            cands = [u for _, u in sorted(weak.get(s, []))][:6] or [f"{s} 総合演習"]
            items = [{"unit": u, "koma": 0} for u in cands]
            for i in range(total):
                items[i % len(items)]["koma"] += 1
            items = [it for it in items if it["koma"] > 0]
        diff = total - sum(it["koma"] for it in items)
        i = 0
        while diff != 0 and items:
            it = items[i % len(items)]
            if diff > 0:
                it["koma"] += 1
                diff -= 1
            elif it["koma"] > 1:
                it["koma"] -= 1
                diff += 1
            i += 1
            if i > 10000:
                break
        out[s] = items
    return out


def _run_mendan(req: MendanReq) -> dict:
    """Core generation (blocking). Returns a result dict, or {"error": ...} on
    failure. Shared by the HTTP POST and the WebSocket endpoints."""
    cc, sc = _code(req.classroom), _code(req.student)
    data = _gather_student_data(cc, sc)
    signals = _behavior_signals(data["ctest"], data["pcs"])
    # Parse ETS (ai_parse_document) and SSS (xlsx) concurrently to minimise latency.
    ets_text, sss_text = "", ""
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=2) as ex:
        ets_fut = ex.submit(_parse_ets_text, str(data["ets"].iloc[0]["file_path"])) if not data["ets"].empty else None
        sss_fut = ex.submit(_parse_sss_text, str(data["sss"].iloc[0]["file_path"])) if not data["sss"].empty else None
        own = set(data["sss"]["file_path"].astype(str)) if not data["sss"].empty else set()
        ref_paths = [p for p in (req.ref_sss or [])[:3]
                     if p.startswith(f"/Volumes/{CATALOG}/{SCHEMA}/sss/") and p not in own]
        ref_loaded = list(ex.map(_load_sss, ref_paths))
        if ets_fut:
            ets_text = ets_fut.result()
        if sss_fut:
            sss_text = sss_fut.result()
    ref_blocks, ref_labels = [], []
    for txt, prof in ref_loaded:
        if not txt:
            continue
        lb = _sss_label(prof)
        ref_labels.append(lb)
        ref_blocks.append(f"＜{lb}＞\n{txt[:1800]}")
    ref_text = "\n\n".join(ref_blocks)
    sources = _field_sources(data, ets_used=bool(ets_text), sss_used=bool(sss_text), ref_labels=ref_labels, req=req)
    used_sources = sorted({x for lst in sources.values() for x in lst if not x.startswith("（")})
    free = req.mode == "free"
    prompt = (
        _build_free_prompt(cc, sc, data, req, signals, ets_text, sss_text, ref_text)
        if free else _build_prompt(cc, sc, data, req, signals, ets_text, sss_text, ref_text)
    )
    try:
        system = ("あなたは塾の面談資料を作成する専門家です。読みやすいMarkdownで回答してください。"
                  if free else "あなたは塾の講習提案書を作成する専門家です。JSON形式で回答してください。")
        raw = _chat_llm(
            [{"role": "system", "content": system}, {"role": "user", "content": prompt}],
            # 行数・字数の型を守らせると思考トークンが増え、8000では本文が途中で切れる
            max_tokens=32000, reasoning_effort="medium",
        )
        meta = {
            "ets_parsed": bool(ets_text), "sss_parsed": bool(sss_text), "ref_used": ref_labels,
            "ets_date": _ets_date(ets_text), "ctest_table": _ctest_series(data["ctest"]),
        }
        if free:
            return {"free_text": raw.strip(), "used_sources": used_sources, **meta}
        start, end = raw.find("{"), raw.rfind("}") + 1
        if start >= 0 and end > start:
            try:
                fields = json.loads(raw[start:end])
                # Guarantee the 学習特性・行動評価 section is never blank when we
                # have grounding data (the LLM occasionally omits it — no temperature control).
                if not (fields.get("ets_eval") or "").strip() and (ets_text or signals):
                    fields["ets_eval"] = _fallback_ets_eval(signals, ets_text)
                koma_units = _fix_koma_units(fields.pop("koma_units", None), req, data)
                fields = {k: (v if isinstance(v, str) else str(v or "")) for k, v in fields.items()}
                return {"fields": fields, "koma_units": koma_units, "sources": sources, **meta}
            except json.JSONDecodeError:
                pass
        return {"raw_response": raw, "sources": sources, **meta}
    except Exception as e:
        log.exception("mendan generate error")
        return {"error": str(e)}


@app.post("/api/mendan/generate")
def mendan_generate(req: MendanReq):
    result = _run_mendan(req)
    if "error" in result:
        raise HTTPException(status_code=500, detail=result["error"])
    return result


class ReviseSection(BaseModel):
    key: str
    label: str
    text: str


class ReviseReq(BaseModel):
    classroom: str
    student: str
    question: str
    answer: str
    sql: Optional[str] = None
    columns: list[str] = []
    rows: list[list[Any]] = []
    sections: list[ReviseSection]


def _rows_text(columns: list[str], rows: list[list[Any]], limit: int = 30) -> str:
    if not columns or not rows:
        return "（結果表なし）"
    lines = [" | ".join(columns)]
    lines += [" | ".join(str(c) for c in r) for r in rows[:limit]]
    if len(rows) > limit:
        lines.append(f"（ほか {len(rows) - limit} 行）")
    return "\n".join(lines)


def _run_revise(req: ReviseReq) -> dict:
    """Ask the LLM which existing sections the Genie finding should change, and
    return before/after per section for the user to approve or reject."""
    _code(req.classroom), _code(req.student)
    by_key = {s.key: s for s in req.sections}
    sec_text = "\n\n".join(f"<section key=\"{s.key}\" label=\"{s.label}\">\n{s.text}\n</section>" for s in req.sections)
    prompt = f"""あなたは塾の保護者面談資料を編集する担当者です。
教室長が Genie（成績データに対する自然言語分析）に追加で質問し、次の分析結果を得ました。
この結果を踏まえて、既存の面談資料のうち **修正すべきセクションだけ** を書き直してください。

【Genie への質問】
{req.question}

【Genie の回答】
{req.answer}

【Genie の結果表（出典: Unity Catalog テーブルへの生成SQL）】
{_rows_text(req.columns, req.rows)}

【既存の面談資料（セクションごと）】
{sec_text}

【編集ルール】
- 分析結果と関係のあるセクションだけを対象にする（通常1〜3セクション）。関係のないセクションは出力しない
- 対象セクションは、既存の内容・文体・構成をできるだけ保ったまま、分析結果を自然に織り込んだ **全文** を返す
- Genie の回答・結果表にない数値や事実を作らない。既存の内容と矛盾する場合は、Genie の結果を正として修正する
- 提案書の型を保つ: 「・」で始まる4行（1行70〜85字）の章は4行のまま、前向きな言い方のまま書き直す
{_PARENT_WORDING}
- 出典ラベルや「Genieによると」といった前置きは書かない
- 各セクションに、何をどう変えたかを1文で reason に書く

次のJSONだけを出力してください:
{{"revisions": [{{"key": "<セクションのkey>", "after": "<修正後の全文>", "reason": "<変更理由1文>"}}]}}"""
    try:
        raw = _chat_llm(
            [{"role": "system", "content": "あなたは塾の面談資料の編集者です。JSONだけで回答してください。"},
             {"role": "user", "content": prompt}],
            max_tokens=16000, reasoning_effort="medium",
        )
        start, end = raw.find("{"), raw.rfind("}") + 1
        data = json.loads(raw[start:end]) if start >= 0 and end > start else {}
        revisions = []
        for rv in data.get("revisions", []):
            key = rv.get("key")
            after = (rv.get("after") or "").strip()
            if key not in by_key or not after or after == by_key[key].text.strip():
                continue
            revisions.append({
                "key": key, "label": by_key[key].label,
                "before": by_key[key].text, "after": after, "reason": rv.get("reason", ""),
            })
        return {"revisions": revisions}
    except Exception as e:  # noqa: BLE001
        log.exception("revise error")
        return {"error": str(e)}


@app.post("/api/mendan/revise")
def mendan_revise(req: ReviseReq):
    result = _run_revise(req)
    if "error" in result:
        raise HTTPException(status_code=500, detail=result["error"])
    return result


async def _ws_run(ws: WebSocket, model, fn):
    """Receive one JSON request, run a blocking fn in a thread, send the result.
    WebSockets bypass the 120s Apps reverse-proxy timeout."""
    import asyncio
    await ws.accept()
    try:
        payload = await ws.receive_json()
        req = model(**payload)
        result = await asyncio.get_event_loop().run_in_executor(None, fn, req)
        await ws.send_json(result)
    except WebSocketDisconnect:
        return
    except Exception as e:  # noqa: BLE001
        log.exception("mendan ws error")
        try:
            await ws.send_json({"error": str(e)})
        except Exception:
            pass
    finally:
        try:
            await ws.close()
        except Exception:
            pass


# ──────────────────────────────────────────────
# Genie Agent mode (tables + files in attached Volumes)
# The classic conversation API only queries tables; file analysis in Volumes
# runs only in Agent mode: POST /api/2.0/genie/agents/{space_id}/responses (SSE).
# ──────────────────────────────────────────────
class AgentReq(BaseModel):
    message: str
    conversation_id: Optional[str] = None


def _run_agent(req: AgentReq, emit) -> dict:
    """Call Genie Agent mode and relay its steps through `emit` as they happen.
    Returns the final answer (Markdown with citation links)."""
    import requests

    cfg = get_workspace_client().config
    body: dict[str, Any] = {"input": [{"type": "message", "role": "user",
                                       "content": [{"type": "input_text", "text": req.message}]}]}
    if req.conversation_id:
        body["conversation_id"] = req.conversation_id
    headers = {"Content-Type": "application/json", "Accept": "text/event-stream", **cfg.authenticate()}
    r = requests.post(f"{cfg.host.rstrip('/')}/api/2.0/genie/agents/{GENIE_SPACE_ID}/responses",
                      json=body, headers=headers, stream=True, timeout=(30, 600))
    if r.status_code >= 400:
        r.encoding = "utf-8"
        raise RuntimeError(f"Genie agent error {r.status_code}: {r.text[:300]}")

    answer, conv_id, sqls, file_searches = "", req.conversation_id, [], 0
    response_id = None
    seen: set[str] = set()  # items arrive as "added" and/or "done"; handle each once
    for raw in r.iter_lines():
        if not raw:
            continue
        line = raw.decode("utf-8", errors="replace")  # SSE has no charset header → decode explicitly
        if not line.startswith("data:"):
            continue
        try:
            ev = json.loads(line[5:])
        except json.JSONDecodeError:
            continue
        etype = ev.get("type")
        if etype in ("response.created", "response.completed", "response.failed"):
            resp = ev.get("response") or {}
            conv_id = resp.get("conversation_id") or conv_id
            response_id = resp.get("id") or response_id
            if etype == "response.failed" or resp.get("status") == "failed":
                raise RuntimeError(f"Genie agent failed: {(resp.get('error') or {}).get('message', resp.get('error'))}")
            if etype == "response.completed" and not answer:
                for it in resp.get("output") or []:
                    if it.get("type") == "message" and it.get("role") == "assistant":
                        answer = "".join(c.get("text", "") for c in it.get("content") or [])
            continue
        if etype not in ("response.output_item.added", "response.output_item.done"):
            continue
        it = ev.get("item") or {}
        ity = it.get("type")
        if ity == "message" and it.get("role") == "assistant" and etype == "response.output_item.done":
            answer = "".join(c.get("text", "") for c in it.get("content") or [])
            continue
        iid = it.get("call_id") or it.get("id")
        # reasoning text is only complete on "done"; tool calls can be handled on whichever comes first
        if ity == "reasoning" and etype != "response.output_item.done":
            continue
        if ity not in ("reasoning", "function_call") or not iid or iid.lower() in seen:
            continue
        seen.add(iid.lower())
        if ity == "reasoning":
            txt = "".join(c.get("text", "") for c in it.get("content") or []).strip()
            if txt:
                emit({"kind": "progress", "icon": "💭", "text": txt[:160]})
        elif ity == "function_call":
            name = it.get("name")
            try:
                args = json.loads(it.get("arguments") or "{}")
            except json.JSONDecodeError:
                args = {}
            if name == "execute_sql":
                title = args.get("title") or "SQLを実行"
                sqls.append({"title": title, "sql": args.get("sql", "")})
                emit({"kind": "progress", "icon": "🧮", "text": f"テーブルを検索：{title}"})
            elif name == "search_context_files":
                file_searches += 1
                q = args.get("query") or args.get("question") or ""
                emit({"kind": "progress", "icon": "📂", "text": f"ファイルを検索（Volume）{('：' + q[:60]) if q else ''}"})
            else:
                emit({"kind": "progress", "icon": "🔧", "text": f"ツール実行：{name}"})

    # The stream does not always report every file search; recount from the
    # conversation items so the "ファイル検索 N 回" badge is accurate.
    if conv_id and response_id:
        try:
            ir = requests.get(f"{cfg.host.rstrip('/')}/api/2.0/genie/agents/{GENIE_SPACE_ID}/conversations/{conv_id}/items",
                              headers={**cfg.authenticate()}, timeout=30)
            ir.encoding = "utf-8"
            items = [x for x in ir.json().get("data", [])
                     if (x.get("metadata") or {}).get("message_id") == response_id and x.get("type") == "function_call"]
            file_searches = max(file_searches, sum(1 for x in items if x.get("name") == "search_context_files"))
        except Exception as e:  # noqa: BLE001
            log.warning("agent item recount failed: %s", e)

    if not answer:
        answer = "回答を取得できませんでした。質問を変えてお試しください。"
    return {"kind": "final", "answer": answer.strip(), "conversation_id": conv_id,
            "sqls": sqls, "file_searches": file_searches}


@app.websocket("/api/genie/agent/ws")
async def genie_agent_ws(ws: WebSocket):
    """Streams Agent mode progress, then the final answer. A WebSocket is
    required: agent answers take ~2–3 min, beyond the 120s Apps proxy limit."""
    import asyncio
    await ws.accept()
    loop = asyncio.get_event_loop()
    queue: asyncio.Queue = asyncio.Queue()
    try:
        req = AgentReq(**(await ws.receive_json()))
        emit = lambda msg: loop.call_soon_threadsafe(queue.put_nowait, msg)  # noqa: E731
        fut = loop.run_in_executor(None, _run_agent, req, emit)
        while True:
            getter = asyncio.ensure_future(queue.get())
            done, _ = await asyncio.wait({getter, fut}, return_when=asyncio.FIRST_COMPLETED)
            if getter in done:
                await ws.send_json(getter.result())
                continue
            getter.cancel()
            break
        while not queue.empty():
            await ws.send_json(queue.get_nowait())
        await ws.send_json(fut.result())
    except WebSocketDisconnect:
        return
    except Exception as e:  # noqa: BLE001
        log.exception("genie agent ws error")
        try:
            await ws.send_json({"kind": "final", "error": str(e)})
        except Exception:
            pass
    finally:
        try:
            await ws.close()
        except Exception:
            pass


@app.websocket("/api/mendan/ws")
async def mendan_ws(ws: WebSocket):
    await _ws_run(ws, MendanReq, _run_mendan)


@app.websocket("/api/mendan/revise/ws")
async def mendan_revise_ws(ws: WebSocket):
    await _ws_run(ws, ReviseReq, _run_revise)


# ──────────────────────────────────────────────
# Terminology / 表記ゆれ detection (用語マスター concern)
# ──────────────────────────────────────────────
_ZEN2HAN = str.maketrans(
    "０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ　",
    "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ ",
)
_KANJI_NUM = {"一": "1", "二": "2", "三": "3", "四": "4", "五": "5",
              "六": "6", "七": "7", "八": "8", "九": "9", "〇": "0"}


def _norm_term(s: str) -> str:
    s = (s or "").translate(_ZEN2HAN)
    for k, v in _KANJI_NUM.items():
        s = s.replace(k, v)
    return re.sub(r"[\s・,，、。．.\-ー―－（）()]", "", s).lower()


@app.get("/api/terms")
def terms(classroom: str):
    """Detect PCS 単元名 表記ゆれ: distinct surface forms that normalise to the
    same key. Directly answers 『同じ内容を別の言葉で表現していないか』."""
    cc = _code(classroom)
    df = run_sql(
        f"SELECT DISTINCT `単元名` name FROM {CATALOG}.{SCHEMA}.pcs_results WHERE `教室コード`={cc}"
    )
    groups: dict[str, set] = {}
    for _, r in df.iterrows():
        name = str(r["name"])
        groups.setdefault(_norm_term(name), set()).add(name)
    variants = [sorted(v) for v in groups.values() if len(v) > 1]
    variants.sort(key=lambda x: -len(x))
    return {
        "distinct_terms": int(len(df)),
        "variant_groups": len(variants),
        "examples": variants[:15],
    }


# ──────────────────────────────────────────────
# Static assets (built React app) — must be mounted LAST
# ──────────────────────────────────────────────
_DIST = os.path.join(os.path.dirname(__file__), "client", "dist")
if os.path.isdir(_DIST):
    app.mount("/", StaticFiles(directory=_DIST, html=True), name="static")
else:
    @app.get("/")
    def _no_build():
        return JSONResponse(
            {"error": "client/dist not found — run `npm run build` in client/"},
            status_code=503,
        )


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("DATABRICKS_APP_PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
